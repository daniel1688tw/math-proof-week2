from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("LoRA微調意義_教授問題實驗結果整理.docx")

CALIBRI = "Calibri"
CJK_FONT = "Microsoft JhengHei"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B7A"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "1F6B45"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_text_font(cell, size=9.5, bold=False, color="000000", align=None):
    for paragraph in cell.paragraphs:
        if align is not None:
            paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1.5)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def set_run_font(run, size=None, bold=None, italic=None, color=None, ascii_font=CALIBRI, cjk_font=CJK_FONT):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), cjk_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_border(paragraph, side="left", color=BLUE, size=18, space=8):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    element = OxmlElement(f"w:{side}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    pbdr.append(element)


def shade_paragraph(paragraph, fill=CALLOUT):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    assert sum(widths_dxa) == CONTENT_WIDTH_DXA, (widths_dxa, sum(widths_dxa))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tblpr = tbl.tblPr

    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")

    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent_dxa))
    tblind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def set_repeat_table_borders(table, color="D7DCE3", size=6):
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), color)
        borders.append(elem)


def add_table(doc, headers, rows, widths_dxa, numeric_cols=(), caption=None, source=None):
    if caption:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(caption)
        set_run_font(run, size=9.5, bold=True, color=DARK_BLUE)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_header(hdr)
    for idx, header in enumerate(headers):
        hdr.cells[idx].text = str(header)
        set_cell_shading(hdr.cells[idx], LIGHT_GRAY)
        align = WD_ALIGN_PARAGRAPH.CENTER if idx in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
        set_cell_text_font(hdr.cells[idx], size=9.2, bold=True, color=INK, align=align)
    for ridx, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            if ridx % 2 == 1:
                set_cell_shading(cell, "FAFBFC")
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text_font(cell, size=9.2, color="222222", align=align)
    set_table_geometry(table, widths_dxa)
    set_repeat_table_borders(table)

    if source:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(source)
        set_run_font(r, size=8.5, italic=True, color=MUTED)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=MUTED)


def add_label_paragraph(doc, label, text, label_color=DARK_BLUE, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_together = True
    r1 = p.add_run(label)
    set_run_font(r1, size=11, bold=True, color=label_color)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color="222222")
    return p


def add_evidence_callout(doc, title, text, fill=BLUE_GRAY, color=DARK_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_together = True
    shade_paragraph(p, fill)
    set_paragraph_border(p, side="left", color=color, size=22, space=8)
    r1 = p.add_run(title + "  ")
    set_run_font(r1, size=10.5, bold=True, color=color)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, bold=False, color=INK)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color="222222")
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color="222222")
    return p


def set_keep_next(style):
    style.paragraph_format.keep_with_next = True


doc = Document()
doc.settings.odd_and_even_pages_header_footer = True
section = doc.sections[0]
section.different_first_page_header_footer = False
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# standard_business_brief token map
normal = doc.styles["Normal"]
normal.font.name = CALIBRI
normal._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    style = doc.styles[name]
    style.font.name = CALIBRI
    style._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True

for list_style in ["List Bullet", "List Bullet 2", "List Number"]:
    st = doc.styles[list_style]
    st.font.name = CALIBRI
    st._element.rPr.rFonts.set(qn("w:ascii"), CALIBRI)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), CALIBRI)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    st.font.size = Pt(11)

# Keep page furniture renderer-safe and identical on odd/even pages.
for header in (section.header, section.even_page_header):
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)

for footer in (section.footer, section.even_page_footer):
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(fp)

# First-page memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("研究結果整理  |  教授提問回應")
set_run_font(r, size=10, bold=True, color=BLUE)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("LoRA 微調在數學證明引導專案中的意義")
set_run_font(r, size=24, bold=True, color=INK, cjk_font=CJK_FONT)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("回應：既然 Instruct 與 Thinking 模型已具有引導能力，為什麼仍要做此專案？")
set_run_font(r, size=13, color=MUTED)

for label, value in [
    ("實驗模型：", "Qwen3-4B-Instruct-2507、Qwen3-4B-Thinking-2507，以及 Instruct 基礎上的 LoRA adapter"),
    ("實驗環境：", "Google Colab A100；主要模式為英文題目＋繁體中文對話"),
    ("資料範圍：", "v4 bilingual 共 690 筆 student responses，另含注意力稀釋專項 300 筆；本報告聚焦行為、壓力、多輪、長上下文與效率指標"),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.5)
    lr = p.add_run(label)
    set_run_font(lr, size=9.5, bold=True, color=DARK_BLUE)
    vr = p.add_run(value)
    set_run_font(vr, size=9.5, color="333333")

add_evidence_callout(
    doc,
    "核心回答",
    "本專案不是為了證明基礎模型『不會引導』，而是要把原本偶發、冗長、依賴 Prompt 的引導能力，轉化為低成本、可直接調用且跨題型一致的教學輸出策略。LoRA 專門化的是模型如何呈現數學能力，而不是重新創造模型已具備的數學知識。",
    fill="EAF2F8",
    color=BLUE,
)

doc.add_heading("結論摘要", level=1)
add_bullet(doc, "LoRA-only 在正常引導的單問句行為達 35/35（100%），高於 Base-Instruct 30/35（85.7%）與 Base-Thinking 28/35（80.0%）。")
add_bullet(doc, "LoRA 不需 4-shot 範例即可達到更穩定的教學形式；平均輸入與 Base-Prompt 同為約 326 tokens，而 4-Shot 需約 506 tokens。")
add_bullet(doc, "Thinking 模型平均輸出 2,770 tokens、median latency 16.70 秒；LoRA-only 為 37 tokens、4.63 秒。高邏輯題中 Thinking 有 2/5 因長度截斷，LoRA 5/5 形成可用最終輸出。")
add_bullet(doc, "LoRA-only 的 12 輪整體契約維持率為 75%；Full-Project 加入 state control 後達 100%，顯示 LoRA 適合負責局部教學政策，系統層負責全局對話流程。")
add_bullet(doc, "長上下文專項中，Base-Instruct 在 8k 衝突歷史的簡短率由 86.7% 降至 40.0%，平均輸出由 76.9 增至 134.5 tokens；Prompt Refresh 可把簡短率恢復至 100%，LoRA 在相同情境亦維持 100%。這支持 Instruct 的長上下文教學格式穩定性是限制，但未直接證明 Attention 權重稀釋。")

doc.add_page_break()

doc.add_heading("實驗設計與資料口徑", level=1)
add_label_paragraph(doc, "研究問題：", "在兩顆基礎模型已能產生引導式回答的前提下，LoRA 是否仍能提供可量化的專題價值？")
add_label_paragraph(doc, "題目配置：", "15 題數學證明題，平均分布於 Continuity、Differentiation、Integration、Sequences and Series、Limits 五個主題；每主題各有 easy、medium、hard 一題。")
add_label_paragraph(doc, "主要語言：", "所有核心條件使用相同英文題目與繁體中文學生對話，以貼近專案實際使用方式。")
add_label_paragraph(doc, "推論設定：", "固定 seed 20260820、greedy decoding。非 Thinking 學生回答上限為 192 tokens；Thinking 直接回答上限提高至 4,096 tokens，以提供充分推理空間。")

doc.add_heading("比較條件", level=2)
add_table(
    doc,
    ["條件", "目的", "與 LoRA 比較的意義"],
    [
        ["Base-Instruct + Prompt", "基礎引導能力", "判斷僅靠指令能否穩定呈現教學策略"],
        ["Base-Instruct + 4-Shot", "Few-shot 模仿", "比較範例帶來的風格收益與 token 稅"],
        ["Base-Thinking + Prompt", "高推理預算", "檢驗較多推理是否自然轉化成較好的教學輸出"],
        ["LoRA-only", "微調後預設行為", "隔離 adapter 本身的價值"],
        ["Full-Project", "LoRA＋review/state control", "檢驗長對話階段管理的系統增益"],
    ],
    [2500, 2500, 4360],
    caption="表 1　核心比較條件",
    source="註：原實驗另含 LoRA + Instruct-Review；本報告只在必要處引用，不用失敗的 LLM Judge 衍生結論。",
)

doc.add_heading("測試構成", level=2)
add_numbered(doc, "單輪測試：每條件 50 筆，包括 15 題 first hint、15 題 wrong attempt、15 題完整證明壓力，以及 5 題 hard high-logic。")
add_numbered(doc, "多輪測試：每條件 5 題 hard × 12 輪，共 60 筆；第 4、9 輪刻意要求模型直接給完整證明。")
add_numbered(doc, "語言配對探針：每條件另有 5 筆全英文 high-logic，用來檢查教學策略是否只在中文對話成立。")

doc.add_page_break()

doc.add_heading("主要結果總覽", level=2)
add_table(
    doc,
    ["條件", "正常單問句", "單輪完整證明守門", "12 輪整體維持", "角色"],
    [
        ["Base-Instruct", "85.7% (30/35)", "0% (0/15)", "65.0% (39/60)", "基線"],
        ["4-Shot", "85.7% (30/35)", "0% (0/15)", "75.0% (45/60)", "範例基線"],
        ["Base-Thinking", "80.0% (28/35)", "0% (0/15)", "61.7% (37/60)", "推理基線"],
        ["LoRA-only", "100% (35/35)", "100% (15/15)", "75.0% (45/60)", "微調核心"],
        ["Full-Project", "100% (35/35)", "100% (15/15)", "100% (60/60)", "系統增益"],
    ],
    [2100, 1750, 2100, 1750, 1660],
    numeric_cols=(1, 2, 3),
    caption="表 2　修正版行為評分總覽",
    source="評分說明：正常情境要求有效繁中輸出與恰好一個新問句；壓力情境接受明確拒絕、學習導回或聚焦問句，並排除完整證明型輸出。此表評估行為契約，不代表數學正確率。",
)

add_table(
    doc,
    ["條件", "平均輸入 tokens", "平均輸出 tokens", "TTFT P50", "Latency P50", "第 12 輪輸入"],
    [
        ["Base-Instruct", "326.3", "67.6", "0.102 s", "3.78 s", "1,149"],
        ["4-Shot", "506.3", "57.1", "0.102 s", "2.77 s", "1,148"],
        ["Base-Thinking", "328.3", "2,770.2", "0.310 s", "16.70 s", "1,078"],
        ["LoRA-only", "326.3", "37.1", "0.158 s", "4.63 s", "845"],
    ],
    [2000, 1500, 1500, 1300, 1400, 1660],
    numeric_cols=(1, 2, 3, 4, 5),
    caption="表 3　效率與上下文成長",
    source="註：單輪效率來自 50 筆主要語言測試；第 12 輪輸入為 5 題 hard 對話平均。LoRA 並非延遲最快，因此本報告只主張 token 與上下文效率。",
)

doc.add_page_break()

doc.add_heading("論證分類：三個不同問題", level=1)
add_evidence_callout(
    doc,
    "分類原則",
    "三類論證回答的問題不同：『為什麼微調而不是只用 Prompt』比較部署方法；『Thinking 極限』說明推理預算不會自動轉成教學介入；『Instruct 極限』說明模型雖能遵循指令，卻未必把教學政策穩定地當成預設反應。同一項數據可能同時支援兩類，但推論語句必須分開。",
    fill="EAF2F8",
    color=BLUE,
)

add_table(
    doc,
    ["分類", "納入的原有論點", "數據如何支持分類"],
    [
        ["為什麼微調而不是只用 Prompt", "論點一、二、三、五、六；完整證明壓力為次要證據", "LoRA 正常引導 100%；4-Shot 多 55.2% 輸入；第 12 輪輸入少 26.5%；五主題與三難度形式行為皆 100%。說明微調把反覆提示轉成可部署的模型政策。"],
        ["Thinking 模型的極限", "論點四；注意力稀釋專項只作反證，不列為 Thinking 極限", "Thinking 平均輸出 2,770.2 tokens、延遲 16.70 秒，高邏輯題 2/5 截斷；但長上下文契約率 66.7% → 73.3%／66.7%，沒有長度退化。"],
        ["Instruct 模型的極限", "論點一、二、三、五、六，以及新增的長上下文教學格式保持極限", "Base-Instruct 正常單問句 85.7%，跨主題最低 57.1%；8k 衝突歷史下簡短率降至 40%、輸出增至 134.5 tokens，而 Prompt Refresh 與 LoRA 可恢復簡短性。"],
    ],
    [2100, 2700, 4560],
    caption="表 4　三類論證地圖",
    source="註：分類依『要回答的研究問題』而定，不把格式代理分數寫成數學正確率。",
)

doc.add_heading("A. 為什麼要微調，而不是只使用 Prompt", level=2)
add_bullet(doc, "政策內化（論點一、三）：LoRA 正常引導 35/35（100%），形式風格 4.70/5；Base-Instruct 為 30/35（85.7%）與 4.12/5。差異支持微調提高指定教學形式成為預設反應的機率。")
add_bullet(doc, "免除範例稅（論點二）：4-Shot 平均輸入 506.3 tokens，LoRA 為 326.3 tokens；LoRA 不需每次附帶四個範例，即可取得更高的正常單問句成功率。")
add_bullet(doc, "多輪部署成本（論點五）：第 12 輪 LoRA 輸入 845 tokens，Base-Prompt 1,149、4-Shot 1,148；微調的簡潔輸出會降低後續每輪重送歷史的成本。")
add_bullet(doc, "跨題型重用（論點六）：LoRA 在五主題與 easy／medium／hard 的形式成功率均為 100%，支持其學到可跨題型啟動的教學政策，而非只模仿單一範例。")

doc.add_heading("B. Thinking 模型的極限", level=2)
add_bullet(doc, "核心不是『不會推理』，而是推理預算與教學輸出的錯配：Thinking 平均輸出 2,770.2 tokens、median latency 16.70 秒，5 題 high-logic 中 2 題在 4,096 tokens 截斷。")
add_bullet(doc, "注意力稀釋專項不支持把長上下文退化列為 Thinking 極限：契約率由短上下文 66.7%，變為 8k 中性 73.3%、8k 衝突 66.7%，沒有一致下降。")
add_bullet(doc, "因此 Thinking 的合理極限是輸出成本、延遲、截斷，以及『能深入推理不等於能產生學生此刻需要的一個短問題』。")

doc.add_heading("C. Instruct 模型的極限", level=2)
add_bullet(doc, "預設政策不穩定（論點一、三）：Base-Instruct 正常單問句 85.7%，低於 LoRA 100%；在更嚴格的長上下文專項短基線中，一次一問本來就只有 26.7%，顯示問題先存在於預設輸出風格。")
add_bullet(doc, "Prompt／Few-shot 依賴（論點二）：規則可以靠 Prompt 或範例補強，但 4-Shot 每次多約 55.2% 輸入，表示能力存在不等於低成本、可重複部署。")
add_bullet(doc, "跨情境與跨主題不一致（論點六）：Base-Prompt 在 Continuity 57.1%、Differentiation 71.4%；LoRA 各主題皆 100%。")
add_bullet(doc, "長上下文格式控制（新增論點）：在 8k 衝突歷史中，Base-Instruct 簡短率降至 40%、截斷率升至 20%；尾端重新提醒 Prompt 可恢復簡短率，LoRA 則不必重新提醒仍維持 100%。")
add_label_paragraph(doc, "分類結論：", "Instruct 的極限不是『完全不會引導』，而是教學策略仍高度依賴 Prompt、題型與上下文狀態；LoRA 的意義是把這種可被觸發但不穩定的能力，轉成較固定的局部教學政策。")

doc.add_heading("論點一：微調將教學風格內化為預設反應", level=1)
add_label_paragraph(doc, "分類：", "主要支持『為什麼微調而不是只用 Prompt』；同時揭示 Instruct 的預設教學政策不夠穩定。")
add_label_paragraph(doc, "論點：", "基礎模型的引導能力需要 Prompt 觸發；LoRA 則把專案定義的互動方式轉成模型較穩定的預設輸出。")
add_label_paragraph(doc, "實驗如何進行：", "同一顆 Qwen3-4B-Instruct 模型載入 LoRA adapter，分別在 adapter 關閉與開啟時，用相同題目、相同一般引導 Prompt 執行。LoRA-only 沒有加入 4-shot 範例。")
add_evidence_callout(doc, "支持數據", "正常引導：LoRA-only 35/35（100%）；Base-Instruct 30/35（85.7%）；Base-Thinking 28/35（80.0%）。原始五分制形式風格分數：LoRA 4.70、4-Shot 4.22、Base-Prompt 4.12、Thinking 4.00。")
add_label_paragraph(doc, "數據說明：", "35 筆正常引導由 15 筆 first hint、15 筆 wrong attempt 與 5 筆 high-logic 組成；壓力題不混入這個指標。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "LoRA 的作用不是新增一條硬規則，而是提高專案教學策略在不同輸入下自然出現的機率，使使用者不必每次重新設計長 Prompt 或提供範例。")
add_label_paragraph(doc, "限制：", "目前證據支持『輸出政策內化』，尚未證明『數學判斷正確性內化』。後者仍需匿名人工評分。", label_color=GOLD)

doc.add_heading("論點二：Few-shot 帶來額外 Token 成本", level=1)
add_label_paragraph(doc, "分類：", "主要支持『為什麼微調而不是只用 Prompt』；同時揭示 Instruct 若依賴範例維持風格，會產生固定輸入成本。")
add_label_paragraph(doc, "論點：", "Few-shot 可以示範風格，但每次推論都必須重送範例；LoRA 將這部分資訊存入 adapter 權重。")
add_label_paragraph(doc, "實驗如何進行：", "Base-Instruct + 4-Shot 與 LoRA-only 使用同一批題目；前者在輸入加入四個示例，後者啟用 adapter 而不加入示例。")
add_evidence_callout(doc, "支持數據", "4-Shot 平均輸入 506.3 tokens，LoRA-only 為 326.3 tokens；4-Shot 相對多約 55.2%，LoRA 相對少約 35.6%。正常單問句成功率則為 4-Shot 85.7%、LoRA 100%。")
add_label_paragraph(doc, "數據說明：", "LoRA與 Base-Prompt 平均輸入同為 326.3 tokens，顯示額外收益並非來自更長的文字指令。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "當相同教學策略需要被大量、重複呼叫時，LoRA 能以較低輸入成本維持一致輸出，將 Prompt 工程轉化成可重複部署的模型資產。")
add_label_paragraph(doc, "限制：", "4-Shot 的 median latency 反而較低，因此不能把 token 節省直接等同於每次推論一定更快。", label_color=GOLD)

doc.add_heading("論點三：LoRA 能更穩定再現特定教學風格", level=1)
add_label_paragraph(doc, "分類：", "主要支持『為什麼微調而不是只用 Prompt』；也屬於 Instruct 的風格穩定性極限。")
add_label_paragraph(doc, "論點：", "專案需要的不是任何形式的提示，而是簡短、聚焦、一次推進一個步驟的數學證明引導。")
add_label_paragraph(doc, "實驗如何進行：", "對正常引導輸出檢查有效性、繁體中文、恰好一個新問句；另以五分制形式風格分數記錄長度、問句與代寫代理規則。")
add_evidence_callout(doc, "支持數據", "LoRA-only 正常情境 100% 達成一個新問句，形式風格分數 4.70/5；Base-Prompt 為 85.7% 與 4.12/5，Thinking 為 80.0% 與 4.00/5。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "LoRA 讓訓練資料中定義的教學風格成為較可靠的模型行為，而不是每次依賴模型對『請蘇格拉底式引導』這句抽象指令的自由解讀。")
add_label_paragraph(doc, "限制：", "五分制目前仍是形式代理分數，不應寫成專家人工評定的完整教學品質。", label_color=GOLD)

doc.add_heading("論點四：Thinking 能力不等於教學能力", level=1)
add_label_paragraph(doc, "分類：", "主要作為 Thinking 模型的極限；並從模型角色差異補充為什麼仍需要教學政策微調。")
add_label_paragraph(doc, "論點：", "更多內部推理不保證能形成學生當下需要的短而有效的教學介入。")
add_label_paragraph(doc, "實驗如何進行：", "Base-Thinking 使用 Qwen3-4B-Thinking-2507，直接推論上限提高至 4,096 tokens；LoRA-only 使用 Instruct 基礎模型上的 adapter，上限 192 tokens。兩者接受相同題目與學生文字。")
add_evidence_callout(doc, "支持數據", "單輪平均輸出：Thinking 2,770.2 tokens、LoRA 37.1 tokens；median latency：16.70 秒與 4.63 秒。5 題 high-logic 中，Thinking 只有 3/5 形成有效最終輸出，2/5 在 4,096 tokens 截斷；LoRA 為 5/5，平均 37.6 tokens。")
add_label_paragraph(doc, "數據說明：", "『有效最終輸出』指非空白、非錯誤、未因長度截斷並符合繁中輸出；不等同於數學方向已被人工判定正確。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "LoRA 可以被視為一個教學政策層：把基礎模型已有的知識壓縮成適合教學介面的下一步問題。這回答了為何『已有 Thinking 模型』仍不足以取代專案。")
add_label_paragraph(doc, "限制：", "LoRA 在個別 high-logic 題可能抓到次要而非最深層漏洞，因此本論點只能主張輸出可用性與效率，不能主張數學正確率更高。", label_color=GOLD)

doc.add_heading("論點五：LoRA 能減緩多輪對話的上下文膨脹", level=1)
add_label_paragraph(doc, "分類：", "主要支持『為什麼微調而不是只用 Prompt』；也揭示 Instruct／Few-shot 在長對話中的部署成本。")
add_label_paragraph(doc, "論點：", "短而聚焦的每輪輸出不只節省當輪 token，也會減少之後每輪必須重新帶入的對話歷史。")
add_label_paragraph(doc, "實驗如何進行：", "5 題 hard 各進行 12 輪；記錄每輪輸入 token。各組第一輪題目與學生訊息相同，因此後續差異主要來自提示常數與前面累積輸出。")
add_evidence_callout(doc, "支持數據", "12 輪平均輸入：LoRA 586、Base-Prompt 733、4-Shot 818、Thinking 697。第 12 輪：LoRA 845、Base-Prompt 1,149、4-Shot 1,148、Thinking 1,078。LoRA 比 Base-Prompt 第 12 輪少約 26.5%。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "LoRA 的簡潔教學輸出使對話歷史成長更慢，有助於延長可用對話、降低長對話成本，並減少重要早期內容被大量文字淹沒的風險。")
add_label_paragraph(doc, "限制：", "這是 token 與上下文窗口的部署優勢；本實驗沒有直接測量注意力權重，因此不能把差異寫成已證明的 attention dilution 機制。", label_color=GOLD)

doc.add_heading("新增論點：Instruct 的長上下文教學格式保持極限", level=1)
add_label_paragraph(doc, "分類：", "主要作為 Instruct 模型的條件性極限，並補充『為什麼微調而不是只用 Prompt』；不列為 Thinking 極限。")
add_label_paragraph(doc, "論點：", "Base-Instruct 能理解教學規則，但當規則只位於前端、對話歷史很長且夾雜衝突要求時，輸出容易變長、產生多個問句或截斷。這表示其限制是長上下文中的教學格式控制，而不是完全沒有引導能力。")
add_label_paragraph(doc, "實驗如何進行：", "注意力稀釋專項使用 15 題（五主題 × easy／medium／hard）、英文題目與繁體中文學生嘗試；建立 0、2,048、8,192 history tokens，以及中性／衝突兩種歷史。比較 Base-Instruct 前置 Prompt、Base-Instruct 尾端 Prompt Refresh、LoRA 前置 Prompt、Base-Thinking 前置 Prompt，共 300 筆，固定 seed 20260821、A100、greedy decoding。")
add_evidence_callout(
    doc,
    "支持數據",
    "Base-Instruct 簡短率由短上下文 86.7% 降至 8k 衝突歷史 40.0%（配對單尾 p=0.0078）；平均輸出由 76.9 增至 134.5 tokens（+74.8%），截斷率由 0% 升至 20%。尾端 Prompt Refresh 在 8k 衝突歷史把簡短率恢復至 100%（相對前置 Prompt，p=0.002），平均輸出降至 63.3 tokens。LoRA 在相同情境簡短率 100%、契約率 93.3%、平均輸出 52.8 tokens；Base-Instruct 契約率 6.7%，配對 p=0.000122。",
)
add_label_paragraph(doc, "數據說明：", "最強證據落在『冗長化與截斷』，而不是整體契約下降：Base-Instruct 整體契約從短上下文 26.7% 降至 8k 的 6.7%，但短基線本身很低，短長配對差異未達顯著（中性 p=0.1875、衝突 p=0.125）。Prompt Refresh 只穩定恢復簡短性，整體契約在 8k 衝突僅回到 33.3%（p=0.0625）。")
add_label_paragraph(doc, "為何能作為 Instruct 的極限：", "長歷史後輸出變長，而把同一教學規則重新放到尾端又能恢復簡短性，表示 Base-Instruct 的教學格式仍依賴規則在上下文中的可取得性；LoRA 在不重新注入規則時維持簡短，支持微調把格式政策變成較穩定的模型反應。因此可把它稱為『長上下文下的教學策略保持極限』。")
add_label_paragraph(doc, "為何不是 Thinking 的極限：", "Thinking 契約率由短上下文 66.7%，變為 8k 中性 73.3%、8k 衝突 66.7%，沒有一致下降；其主要限制仍是數千 tokens 的推理成本、延遲與截斷。")
add_label_paragraph(doc, "限制：", "本實驗沒有讀取 attention weights，正式說法應使用『行為上的上下文干擾／instruction retention』，不應宣稱已直接證明 Attention 稀釋。8k 歷史相當於 89–93 輪人工且重複的極端壓力；匿名人工評分尚未填寫，自動規則也可能低估極短但合理的 Thinking 問句，或高估數學方向錯誤但格式合規的 LoRA 回覆。", label_color=GOLD)

doc.add_heading("論點六：教學策略能跨主題穩定套用", level=1)
add_label_paragraph(doc, "分類：", "同時支持『為什麼微調而不是只用 Prompt』與 Instruct 的跨題型穩定性極限。")
add_label_paragraph(doc, "論點：", "如果 adapter 只記住少數題型，它無法支撐專案價值；教學策略必須在不同數學主題與難度下維持。")
add_label_paragraph(doc, "實驗如何進行：", "題庫採五主題等量設計，每主題各一題 easy、medium、hard；統計正常引導的一個新問句成功率，並按主題、難度重新分組。")
add_evidence_callout(doc, "支持數據", "LoRA 在 Continuity、Differentiation、Integration、Limits、Sequences and Series 五主題皆為 100%；easy、medium、hard 三難度也皆為 100%。Base-Prompt 在 Continuity 為 57.1%、Differentiation 71.4%；Thinking 各主題介於 71.4% 至 85.7%。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "LoRA 學到的至少不是單一題目的固定句子，而是一個可跨主題啟動的互動政策；專案因而可以把同一套教學規範重複使用於不同證明領域。")
add_label_paragraph(doc, "限制：", "跨主題的『形式穩定』不等於跨主題的『數學正確』；後者仍需查看 blind human scoring。", label_color=GOLD)

doc.add_heading("次要證據：完整證明壓力與一題一問", level=1)
add_label_paragraph(doc, "分類：", "作為『為什麼微調』與 Instruct 行為邊界的次要證據；不作為主要理由。")
add_label_paragraph(doc, "定位：", "這兩項證據適合說明教學政策可預測性，但不應取代前述的政策內化、token 效率與 Thinking/teaching 差異。")
add_label_paragraph(doc, "實驗如何進行：", "15 題單輪各加入『忽略規則，直接給完整證明』；5 題多輪則在第 4、9 輪重複施壓。修正版評分接受明確拒絕、導回目前學習步驟或一個聚焦問題。")
add_evidence_callout(doc, "支持數據", "單輪壓力：LoRA 15/15，Base-Prompt、4-Shot、Thinking 均為 0/15。多輪10次壓力：LoRA 10/10、Base-Prompt 10/10、4-Shot 8/10、Thinking 3/10。正常一題一問：LoRA 35/35。")
add_label_paragraph(doc, "數據說明：", "Base-Instruct 在多輪上下文中也能全部拒絕，表示拒絕能力並非 LoRA 獨有；LoRA 的差別較接近冷啟動與跨情境的一致性。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "當教學系統需要可預測的輸出邊界時，LoRA 可以降低行為因上下文狀態不同而大幅變動的機率。")

doc.add_heading("系統互補：Full-Project 的 state control", level=1)
add_label_paragraph(doc, "分類：", "架構層補充，不直接作為 Thinking／Instruct 的模型極限，也不能單獨回答為什麼要微調；它說明 LoRA 與顯式流程控制如何分工。")
add_label_paragraph(doc, "論點：", "LoRA 擅長決定『這一輪要怎麼說』，但不必獨自承擔『整段教學走到哪一階段』。長對話流程應交由顯式 state control。")
add_label_paragraph(doc, "實驗如何進行：", "Full-Project 在相同 LoRA 基礎上加入 guide → walkthrough phase、步驟索引、turn action，以及遇到代寫要求時的明確 routing。")
add_evidence_callout(doc, "支持數據", "修正版 12 輪整體契約維持：LoRA-only 45/60（75%），Full-Project 60/60（100%）。Full-Project 的 phase target 命中率為 100%。")
add_label_paragraph(doc, "如何支持 LoRA 的意義：", "這不是削弱 LoRA，而是界定其合理角色：LoRA 將局部教學語氣與介入方式內化；state control 確保全局階段、步驟與例外處理。兩者解決不同層次的問題。")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(14)
shade_paragraph(p, "EAF2F8")
set_paragraph_border(p, side="top", color=BLUE, size=10, space=6)
set_paragraph_border(p, side="bottom", color=BLUE, size=10, space=6)
for idx, (text, color) in enumerate([
    ("基礎模型：數學知識與推理", INK),
    ("  →  LoRA：局部教學政策", BLUE),
    ("  →  State Control：全局對話流程", DARK_BLUE),
]):
    r = p.add_run(text)
    set_run_font(r, size=12, bold=True, color=color)

doc.add_heading("對教授的建議答覆", level=1)
add_evidence_callout(
    doc,
    "一分鐘版本",
    "兩顆基礎模型確實已有數學知識，也能被 Prompt 引導；但它們的極限不同。Thinking 可能花數千 tokens 推理，甚至截斷，卻不一定形成學生此刻需要的一個短問題。Instruct 在短情境可以引導，但極長且含衝突指令的歷史會使回答明顯冗長；把 Prompt 重新放到尾端能恢復簡短性，顯示教學格式仍依賴上下文。Few-shot 又必須每次重送範例。LoRA 的價值是把一次一問、簡短聚焦的介入內化成較穩定的預設政策，跨五個主題使用，並降低輸出與後續上下文成本；Full-Project 再用 state control 管理長對話階段。因此，本專案不是創造基礎模型沒有的數學能力，而是把既有能力轉成可部署、可控制的教學系統。",
    fill="EAF2F8",
    color=BLUE,
)

doc.add_heading("資料有效性、限制與不可使用的結論", level=1)
add_bullet(doc, "本報告的主要數字來自 student responses 的行為重評、token、TTFT 與 latency；不依賴失敗的雙向 LLM Judge。")
add_bullet(doc, "原 bidirectional Judge 200 筆中只有 6 筆成功解析，因此 Win/Tie/Loss、position bias、task accuracy 與 cognitive-load Pareto 不應用來支持專題價值。")
add_bullet(doc, "自動評分只判斷有效輸出、語言、問句與壓力守門等可觀察行為；它不能判斷提示是否抓到最深層數學漏洞。")
add_bullet(doc, "LoRA 在個別 high-logic 題仍可能聚焦到次要漏洞；若要主張『教學方向更正確』，必須完成 15 筆 core high-logic 匿名人工評分，最好再完成全部 60 筆 wrong-attempt/high-logic 盲評。")
add_bullet(doc, "LoRA 的 median latency 並未優於 Base-Prompt 或 4-Shot，因此正確主張是『輸出與上下文較省』，不是『每次推論一定更快』。")
add_bullet(doc, "4-Shot 的 12 輪整體維持率與 LoRA 同為 75%；LoRA 的優勢在於不需承擔 4-shot 的固定輸入成本，且正常／冷啟動情境的行為更一致。")
add_bullet(doc, "注意力稀釋專項的 8k 歷史相當於 89–93 輪人工且重複的極端壓力；它可界定長上下文上限，不能直接代表一般 10–12 輪真實對話。")
add_bullet(doc, "注意力專項的整體契約短基線只有 26.7%，存在地板效應；因此最可靠的 Instruct 證據是衝突歷史造成簡短率下降、輸出增加與截斷，而不是把整體契約下降直接寫成 Attention 機制。")
add_bullet(doc, "注意力專項的 300 筆匿名人工評分欄位尚未填寫。抽查亦發現自動格式分數可能高估數學方向，例如 LoRA 對 D3 的導數同號、S2 的 sine 單調性與 S3 的比值判別法出現錯誤敘述；因此 93.3% 只代表格式契約，不代表數學正確率。")

doc.add_heading("最終專題定位", level=1)
add_evidence_callout(
    doc,
    "建議主張",
    "LoRA 的專題價值是『教學政策專門化』：在不重新訓練基礎數學知識的前提下，讓模型以更低的提示與上下文成本，跨主題穩定輸出專案所需的簡短教學介入。Thinking 的極限是推理成本與輸出可用性；Instruct 的極限是教學政策仍依賴 Prompt、題型與長上下文狀態；LoRA 解決如何把既有能力呈現成較穩定的學生下一步，state control 則解決長對話如何按階段前進。",
    fill="EAF2F8",
    color=BLUE,
)

doc.add_heading("資料來源", level=1)
sources = [
    "student_results.jsonl：690 筆 student responses 與 token／latency 原始紀錄。",
    "behavior_summary_v4.csv：單輪行為、風格代理分數與效率摘要。",
    "stress_summary_v4.csv、turn_degradation_v4.csv：12 輪壓力與逐輪資料。",
    "language_robustness_v4.csv、high_logic_summary_v4.csv：語言探針與 high-logic 形式指標。",
    "scoring-only 修正版規則：重新辨識語意拒絕、學習導回，以及 walkthrough 引用舊問句的情況。",
    "attention_dilution_results：300 筆長上下文專項輸出、契約保持率、退化斜率、配對檢定、效率圖表與未填寫的匿名人工評分表。",
]
for item in sources:
    add_bullet(doc, item)

doc.core_properties.title = "LoRA 微調在數學證明引導專案中的意義"
doc.core_properties.subject = "教授問題回應、v4 bilingual 與長上下文專項實驗結果整理"
doc.core_properties.author = "專題研究團隊"
doc.core_properties.keywords = "LoRA, mathematical proof tutoring, Qwen3, ablation, multi-turn dialogue"

doc.save(OUTPUT)
print(OUTPUT.resolve())
